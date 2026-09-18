"""Word (.docx) summary report generation (F8)."""

import io

from docx import Document


def generate_summary_report(summary):
    """
    Build a .docx summary of overall risk posture and roadmap status from a
    precomputed `summary` dict — see api/views/reports.py for its shape.
    Returns an in-memory BytesIO positioned at the start of the file.
    """
    document = Document()
    document.add_heading('Risk & Control Roadmap — Summary Report', level=1)

    document.add_heading('Overall Risk Posture', level=2)
    document.add_paragraph(
        'Total residual annualised loss exposure: ${:,.0f}'.format(
            summary.get('totalResidualAle', 0)))
    risk_counts = summary.get('riskCounts', {})
    document.add_paragraph(
        'Risks — open: {}, in progress: {}, closed: {}'.format(
            risk_counts.get('open', 0),
            risk_counts.get('in_progress', 0),
            risk_counts.get('closed', 0)))

    document.add_heading('Top Coverage Gaps', level=2)
    top_gaps = summary.get('topGaps', [])
    if top_gaps:
        table = document.add_table(rows=1, cols=2)
        table.style = 'Table Grid'
        table.rows[0].cells[0].text = 'Technique'
        table.rows[0].cells[1].text = 'Mapped controls'
        for gap in top_gaps:
            row = table.add_row().cells
            row[0].text = '{} ({})'.format(
                gap.get('techniqueName', ''), gap.get('techniqueId', ''))
            row[1].text = str(gap.get('controlCount', 0))
    else:
        document.add_paragraph('No coverage gaps identified.')

    document.add_heading('Roadmap', level=2)
    roadmap = summary.get('roadmap', [])
    if roadmap:
        table = document.add_table(rows=1, cols=4)
        table.style = 'Table Grid'
        for index, title in enumerate(['Control', 'Quarter', 'Status', 'Rationale']):
            table.rows[0].cells[index].text = title
        for item in roadmap:
            row = table.add_row().cells
            row[0].text = item.get('controlName', '')
            row[1].text = item.get('quarter', '')
            row[2].text = item.get('status', '')
            row[3].text = item.get('rationale', '')
    else:
        document.add_paragraph('No roadmap items planned yet.')

    buffer = io.BytesIO()
    document.save(buffer)
    buffer.seek(0)
    return buffer
